from openpyxl.chart import LineChart, Reference, BarChart


def add_chart(
        worksheet,
        chart_type,
        data_columns,
        category_column,
        start_row,
        end_row,
        chart_position,
        chart_title,
        x_axis_title,
        y_axis_title,
        chart_width=25,  # width in cm
        chart_height=15  # height in cm
):
    """
    Add a chart to an Excel worksheet.

    Parameters:
    -----------
    worksheet : openpyxl.Worksheet
        The worksheet where the chart will be added.
    chart_type : str
        Type of chart: 'line' or 'bar'.
    data_columns : int
        Column number containing the data series.
    category_column : int
        Column number containing the categories (X-axis labels).
    start_row : int
        Starting row of the data.
    end_row : int
        Ending row of the data.
    chart_position : str
        Top-left cell where the chart will be placed (e.g., "A12").
    chart_title : str
        Title of the chart.
    x_axis_title : str
        Title for the X-axis.
    y_axis_title : str
        Title for the Y-axis.
    chart_width : float, optional
        Width of the chart in centimeters (default: 25).
    chart_height : float, optional
        Height of the chart in centimeters (default: 15).

    Notes:
    ------
    - For line charts, the line color and width are set by default.
    - `data_columns` should be a single column index (int).
    - Supports basic line and bar charts.
    - Automatically sets categories and titles from the worksheet data.

    Example:
    --------
    add_chart(
        worksheet=worksheet,
        chart_type='line',
        data_columns=7,      # Weight column
        category_column=2,   # Names column
        start_row=7,
        end_row=20,
        chart_position="A12",
        chart_title="Weight Changes by Warehouse",
        x_axis_title="Warehouses",
        y_axis_title="Weight (kg)"
    )
    """

    if chart_type == 'line':
        chart = LineChart()
        chart.style = 20
    elif chart_type == 'bar':
        chart = BarChart()
    else:
        raise ValueError("chart_type must be 'line' or 'bar'.")

    chart.title = chart_title
    chart.y_axis.title = y_axis_title
    chart.x_axis.title = x_axis_title
    chart.width = chart_width
    chart.height = chart_height

    categories = Reference(worksheet, min_col=category_column, min_row=start_row, max_row=end_row)
    data = Reference(worksheet, min_col=data_columns, min_row=start_row - 1, max_row=end_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    # Customize line style for line charts
    for series in chart.series:
        if chart_type == 'line':
            series.graphicalProperties.line.solidFill = "277358"  # default line color
            series.graphicalProperties.line.width = 30000

    worksheet.add_chart(chart, chart_position)
